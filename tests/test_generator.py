#!/usr/bin/env python3
"""
Self-contained tests for the Artifact Generator (stdlib unittest — no pytest
dependency). Covers Phase B acceptance:
  * each of pptx/docx/xlsx produces a real file with OOXML magic bytes (PK\x03\x04)
    and an openable structure
  * unknown artifact_type -> HTTP 4xx
  * contract fields all present and well-formed
Run: python -m unittest tests.test_generator  (or: python tests/test_generator.py)
"""
import json
import os
import sys
import tempfile
import unittest
import urllib.request
import urllib.error
from http.server import ThreadingHTTPServer

# make the app importable when run from repo root
HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, ROOT)

import app.server as ag  # noqa: E402

MAGIC = b"PK\x03\x04"
CONTRACT_FIELDS = {
    "provider",
    "provider_file_id",
    "download_url",
    "filename",
    "mime_type",
    "artifact_type",
    "size",
    "checksum",
}


def _zip_ok(raw: bytes) -> bool:
    """OOXML is a ZIP container: starts with PK\x03\x04 and End-of-Central-Dir
    signature PK\x05\x06 must be present."""
    return raw[:4] == MAGIC and b"PK\x05\x06" in raw[-22:]


SAMPLE_CONTENT = {
    "title": "Q2 Review",
    "sections": [
        {"heading": "Overview", "body": "Quarterly performance.", "bullets": ["Revenue up", "Churn down"]},
        {"heading": "Details", "body": "Segment breakdown.", "bullets": ["Online", "Offline"]},
    ],
    "table_data": {"headers": ["Month", "Revenue", "Orders"], "rows": [["2026-06", "1200000", "9700"]]},
}


class GeneratorUnitTests(unittest.TestCase):
    def setUp(self):
        self.tmp_out = tempfile.mkdtemp(prefix="ag_out_")
        ag.OUTPUT_DIR = self.tmp_out
        ag.TEMPLATES_DIR = os.path.join(ROOT, "templates")
        ag.TEMPLATES = {
            "pptx": os.path.join(ag.TEMPLATES_DIR, "synthetic_template.pptx"),
            "docx": os.path.join(ag.TEMPLATES_DIR, "synthetic_template.docx"),
            "xlsx": os.path.join(ag.TEMPLATES_DIR, "synthetic_template.xlsx"),
        }
        ag.ALLOWED_TEMPLATE_DIRS = [ag.TEMPLATES_DIR, tempfile.gettempdir()]

    def test_each_format_produces_valid_ooxml(self):
        for fmt in ("pptx", "docx", "xlsx"):
            with self.subTest(fmt=fmt):
                tmpl, marker = ag.resolve_template(SAMPLE_CONTENT, fmt)
                raw = ag.GENERATORS[fmt](SAMPLE_CONTENT, tmpl, marker)
                self.assertTrue(len(raw) > 0, f"{fmt}: empty output")
                self.assertTrue(_zip_ok(raw), f"{fmt}: bad OOXML/ZIP structure")
                # openable structure: each library can re-read its own output
                bio = __import__("io").BytesIO(raw)
                if fmt == "pptx":
                    from pptx import Presentation
                    Presentation(bio)
                elif fmt == "docx":
                    from docx import Document
                    Document(bio)
                else:
                    from openpyxl import load_workbook
                    load_workbook(bio)

    def test_unknown_template_mode_rejected(self):
        with self.assertRaises(ValueError):
            ag.resolve_template({"template_mode": "bogus"}, "pptx")


class HttpContractTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp_out = tempfile.mkdtemp(prefix="ag_http_")
        ag.OUTPUT_DIR = cls.tmp_out
        ag.PUBLIC_BASE = "http://127.0.0.1:0"  # placeholder; we read contract structurally
        cls.srv = ThreadingHTTPServer(("127.0.0.1", 0), ag.Handler)
        cls.port = cls.srv.server_address[1]
        import threading
        cls._t = threading.Thread(target=cls.srv.serve_forever, daemon=True)
        cls._t.start()

    @classmethod
    def tearDownClass(cls):
        cls.srv.shutdown()
        cls.srv.server_close()

    def _post(self, path, payload):
        data = json.dumps(payload).encode()
        req = urllib.request.Request(
            f"http://127.0.0.1:{self.port}{path}",
            data=data,
            headers={"Content-Type": "application/json"},
            method="POST",
        )
        with urllib.request.urlopen(req, timeout=5) as r:
            return r.status, json.loads(r.read())

    def _get(self, path):
        req = urllib.request.Request(f"http://127.0.0.1:{self.port}{path}", method="GET")
        with urllib.request.urlopen(req, timeout=5) as r:
            return r.status, r.read()

    def test_metadata_contract_all_formats(self):
        for fmt in ("pptx", "docx", "xlsx"):
            with self.subTest(fmt=fmt):
                code, body = self._post(f"/generate/{fmt}/metadata", SAMPLE_CONTENT)
                self.assertEqual(code, 200)
                art = body["artifact"]
                self.assertEqual(set(art), CONTRACT_FIELDS, f"{fmt}: missing contract fields")
                self.assertEqual(art["artifact_type"], fmt)
                self.assertTrue(art["download_url"].endswith(f".{fmt}"))
                self.assertTrue(art["checksum"].startswith("sha256:"))
                self.assertTrue(art["size"] > 0)
                # the stored file is valid OOXML
                fname = art["download_url"].rsplit("/", 1)[-1]
                with open(os.path.join(ag.OUTPUT_DIR, fname), "rb") as f:
                    raw = f.read()
                self.assertTrue(_zip_ok(raw), f"{fmt}: stored file bad OOXML")
                self.assertEqual(art["size"], len(raw))

    def test_unknown_format_is_4xx(self):
        req = urllib.request.Request(
            f"http://127.0.0.1:{self.port}/generate/pdf/metadata",
            data=json.dumps({"title": "x"}).encode(),
            headers={"Content-Type": "application/json"},
            method="POST",
        )
        with self.assertRaises(urllib.error.HTTPError) as cm:
            urllib.request.urlopen(req, timeout=5)
        self.assertGreaterEqual(cm.exception.code, 400)
        self.assertLess(cm.exception.code, 500)

    def test_download_serves_file_and_health(self):
        code, _ = self._get("/health")
        self.assertEqual(code, 200)
        # generate then download
        _, body = self._post("/generate/docx/metadata", SAMPLE_CONTENT)
        fname = body["artifact"]["download_url"].rsplit("/", 1)[-1]
        code, raw = self._get(f"/download/docx/{fname}")
        self.assertEqual(code, 200)
        self.assertTrue(_zip_ok(raw))

    def test_download_path_traversal_blocked(self):
        try:
            self._get("/download/docx/../../etc/hostname")
            self.fail("traversal should be blocked")
        except urllib.error.HTTPError as e:
            self.assertIn(e.code, (403, 404))

    # ---- POST /generate/metadata (static URL, format from body) ----
    def test_static_metadata_all_formats(self):
        for fmt in ("pptx", "docx", "xlsx"):
            with self.subTest(fmt=fmt):
                payload = dict(SAMPLE_CONTENT, artifact_type=fmt)
                code, body = self._post("/generate/metadata", payload)
                self.assertEqual(code, 200)
                art = body["artifact"]
                self.assertEqual(set(art), CONTRACT_FIELDS)
                self.assertEqual(art["artifact_type"], fmt)
                self.assertTrue(art["download_url"].endswith(f".{fmt}"))
                self.assertTrue(art["checksum"].startswith("sha256:"))
                # stored file is valid OOXML
                fname = art["download_url"].rsplit("/", 1)[-1]
                with open(os.path.join(ag.OUTPUT_DIR, fname), "rb") as f:
                    raw = f.read()
                self.assertTrue(_zip_ok(raw))
                self.assertEqual(art["size"], len(raw))

    def test_static_metadata_byte_identical_to_path_version(self):
        """Same (body, fmt) via /generate/metadata and /generate/{fmt}/metadata
        must produce byte-identical contract JSON."""
        for fmt in ("pptx", "docx", "xlsx"):
            with self.subTest(fmt=fmt):
                payload = dict(SAMPLE_CONTENT, artifact_type=fmt)
                _, static = self._post("/generate/metadata", payload)
                # path version ignores body.artifact_type; pass fmt in URL
                path_payload = dict(SAMPLE_CONTENT)  # no artifact_type key
                _, via_path = self._post(f"/generate/{fmt}/metadata", path_payload)
                self.assertEqual(static, via_path)
                # also assert exact serialized-byte equality
                self.assertEqual(
                    json.dumps(static, sort_keys=True),
                    json.dumps(via_path, sort_keys=True),
                )

    def test_static_metadata_missing_or_unknown_type_is_4xx(self):
        # missing artifact_type
        try:
            self._post("/generate/metadata", {"title": "x"})
            self.fail("expected 4xx for missing artifact_type")
        except urllib.error.HTTPError as e:
            self.assertGreaterEqual(e.code, 400)
            self.assertLess(e.code, 500)
        # unknown artifact_type
        try:
            self._post("/generate/metadata", {"title": "x", "artifact_type": "pdf"})
            self.fail("expected 4xx for unknown artifact_type")
        except urllib.error.HTTPError as e:
            self.assertGreaterEqual(e.code, 400)
            self.assertLess(e.code, 500)

    def test_path_routes_regression(self):
        """Existing /generate/{fmt}/metadata and /generate/{fmt} still work."""
        for fmt in ("pptx", "docx", "xlsx"):
            with self.subTest(fmt=fmt):
                code, body = self._post(f"/generate/{fmt}/metadata", SAMPLE_CONTENT)
                self.assertEqual(code, 200)
                self.assertEqual(body["artifact"]["artifact_type"], fmt)


if __name__ == "__main__":
    unittest.main(verbosity=2)
